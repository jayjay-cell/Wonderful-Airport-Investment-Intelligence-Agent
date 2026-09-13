"""FAA Commercial Service Enplanements and Hub Classification connector.

Static XLSX download, no auth, verified live. Annual data, published with a
~9-month lag after calendar year-end. The workbook already includes an
FAA-assigned Hub code column (L/M/S/N) — no need to compute hub class from
enplanement-share thresholds ourselves.

The whole small workbook (~150KB, ~2000 airports) is downloaded once and
cached in memory; this is NOT the "nationwide dataset" restriction that
applies to T-100 — enplanements is a small annual summary file, not a
per-flight dataset, so downloading it whole is the intended usage pattern
(same as any other small reference table).
"""

from __future__ import annotations

import io
from datetime import date

import httpx
import openpyxl

from core.models import HubClass
from data.cache import cache_get, cache_set
from data.degradation import SourceUnavailable, with_retry

_URL = (
    "https://www.faa.gov/airports/planning_capacity/passenger_allcargo_stats/"
    "passenger/ARP-cy2024-all-enplanements.xlsx"
)
_COVERAGE_YEAR = "CY2024"
_TIMEOUT = 8.0

_HUB_CODE_MAP = {
    "L": HubClass.LARGE,
    "M": HubClass.MEDIUM,
    "S": HubClass.SMALL,
    "N": HubClass.NONHUB,
}

_CACHE_KEY = "faa_enplanements:full_table"


def _load_table() -> dict[str, dict]:
    cached = cache_get(_CACHE_KEY)
    if cached is not None:
        return cached

    def _fetch():
        resp = httpx.get(_URL, timeout=_TIMEOUT, follow_redirects=True)
        resp.raise_for_status()
        return resp.content

    try:
        content = with_retry(_fetch, retries=0)
    except Exception as exc:
        raise SourceUnavailable(
            source_name="FAA Commercial Service Enplanements",
            detail=f"Failed to download enplanements workbook: {exc}",
        ) from exc

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb[wb.sheetnames[0]]

    table: dict[str, dict] = {}
    rows = ws.iter_rows(values_only=True)
    next(rows)  # header row
    for row in rows:
        if not row or len(row) < 11:
            continue
        _, _, state, locid, city, name, _, hub_code, cy_curr, cy_prev, pct_change = row[:11]
        if not locid or not isinstance(locid, str):
            continue
        table[locid.strip().upper()] = {
            "code": locid.strip().upper(),
            "name": name,
            "city": city,
            "state": state,
            "hub_class": _HUB_CODE_MAP.get(hub_code, HubClass.UNKNOWN),
            "enplanements_current_year": cy_curr,
            "enplanements_prior_year": cy_prev,
            "pct_change": pct_change,
        }

    cache_set(_CACHE_KEY, table, ttl_seconds=24 * 60 * 60)
    return table


def get_enplanements(airport_code: str) -> dict | None:
    """Returns the enplanements record for one airport, or None if the code
    is not in the FAA commercial-service enplanements table (e.g. a
    non-commercial or very-low-volume facility)."""
    table = _load_table()
    return table.get(airport_code.strip().upper())


def source_metadata() -> dict:
    return {
        "source_name": "FAA Commercial Service Enplanements and Hub Classification",
        "coverage_period": _COVERAGE_YEAR,
        "retrieved_at": date.today().isoformat(),
    }
